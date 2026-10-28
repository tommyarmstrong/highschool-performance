import argparse
import json
from pathlib import Path
from typing import Literal, NotRequired, TypedDict, cast

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.presentation import Presentation as PresentationDocument
from pptx.shapes.placeholder import SlidePlaceholder
from pptx.oxml.ns import qn
from pptx.oxml.text import CT_TextParagraphProperties
from pptx.oxml.xmlchemy import BaseOxmlElement, OxmlElement
from pptx.slide import Slide, SlideLayout
from pptx.text.text import TextFrame, _Paragraph
from pptx.util import Emu, Inches, Length, Pt

from config.config import get_config
from shared_funcitons import logger


class TitleShape(TypedDict):
    type: Literal["title"]
    text: str
    size: NotRequired[float]


class TextRun(TypedDict):
    text: str
    bold: NotRequired[bool]
    italic: NotRequired[bool]
    hyperlink: NotRequired[str]
    size: NotRequired[float]


class TextParagraph(TypedDict):
    runs: list[TextRun]
    bullet: NotRequired[bool]
    level: NotRequired[int]
    size: NotRequired[float]


class TextBoxShape(TypedDict):
    type: Literal["text_box"]
    paragraphs: list[TextParagraph]
    size: NotRequired[float]
    left: NotRequired[float]
    top: NotRequired[float]
    width: NotRequired[float]
    height: NotRequired[float]


class PictureShape(TypedDict):
    type: Literal["picture"]
    file: str
    left: NotRequired[float]
    top: NotRequired[float]
    width: NotRequired[float]
    height: NotRequired[float]


SlideShape = TitleShape | TextBoxShape | PictureShape


class _Positioned(TypedDict):
    left: float
    top: float
    width: float
    height: float


class SlideDefinition(TypedDict):
    """One slide to build from a template layout and insert at `number`."""

    number: int
    layout: str
    shapes: list[SlideShape]


DEFAULT_DECK = Path("src") / "decks_layout" / "report.json"
_POSITION_KEYS = ("left", "top", "width", "height")

PICTURE_BOTTOM_GAP_INCHES = 0.5


def load_slides(path: Path) -> list[SlideDefinition]:
    """Load a slide deck from JSON and check that each slide can be built."""
    if not path.exists():
        raise FileNotFoundError(f"Slide deck {path} does not exist.")

    with path.open() as deck_file:
        raw = json.load(deck_file)
    if not isinstance(raw, list):
        raise ValueError(f"{path} must contain a list of slides.")

    slides: list[SlideDefinition] = []
    for entry, item in enumerate(raw, start=1):
        slides.extend(_parse_deck_entry(item, entry, path))
    numbers = [slide["number"] for slide in slides]
    duplicates = sorted({number for number in numbers if numbers.count(number) > 1})
    if duplicates:
        raise ValueError(f"{path.name} repeats slide numbers: {duplicates}")
    return slides


def insert_slides(
    prs: PresentationDocument,
    definitions: list[SlideDefinition],
    charts_dir: Path,
) -> None:
    """Build every slide, then put each one at its slide number.

    Slide numbers are 1-based. A listed slide replaces the template slide at
    that number. Numbers past the end append, with no gaps.

    Every slide is added before any template slide is deleted. python-pptx
    names a new slide from the current slide count, so deleting between adds
    makes the next slide reuse the previous new slide's file.
    """
    if not definitions:
        return

    original = list(prs.slides)
    numbers = [definition["number"] for definition in definitions]
    _check_slide_numbers(numbers, len(original))
    spacing = _template_body_spacing(original)

    replaced = {
        definition["number"]: original[definition["number"] - 1]
        for definition in definitions
        if definition["number"] <= len(original)
    }
    created: dict[int, Slide] = {}
    for definition in definitions:
        created[definition["number"]] = _build_slide(prs, definition, charts_dir, spacing)
        logger.info(f"Inserted slide {definition['number']} using {definition['layout']}")

    for original_slide in replaced.values():
        _delete_slide(prs, prs.slides.index(original_slide))

    kept = [slide for slide in original if slide not in replaced.values()]
    kept_slides = iter(kept)
    order = [
        created[number] if number in created else next(kept_slides)
        for number in range(1, len(kept) + len(created) + 1)
    ]
    _reorder_slides(prs, order)


def _build_slide(
    prs: PresentationDocument,
    definition: SlideDefinition,
    charts_dir: Path,
    spacing: tuple[int, int] | None,
) -> Slide:
    """Create a slide from a layout and add its shapes."""
    slide_height = prs.slide_height
    if slide_height is None:
        raise ValueError("Presentation has no slide height")
    slide = prs.slides.add_slide(_layout(prs, definition["layout"]))
    _add_shapes(slide, definition["shapes"], charts_dir, slide_height, spacing)
    return slide


def _check_slide_numbers(numbers: list[int], original_count: int) -> None:
    """Require unique numbers that replace an existing slide or append without gaps."""
    if len(numbers) != len(set(numbers)):
        raise ValueError("Slide numbers must be unique")
    if any(number < 1 for number in numbers):
        raise ValueError("Slide number must be 1 or greater")
    appends = sorted(number for number in numbers if number > original_count)
    expected = list(range(original_count + 1, original_count + len(appends) + 1))
    if appends != expected:
        raise ValueError(
            "Slide numbers must replace an existing slide or append through "
            f"slide {original_count + len(appends)}"
        )


def _layout(prs: PresentationDocument, name: str) -> SlideLayout:
    """Return the slide layout with this name."""
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    names = ", ".join(layout.name for layout in prs.slide_layouts)
    raise ValueError(f"No slide layout named {name!r}. Layouts: {names}")


def _add_shapes(
    slide: Slide,
    shapes: list[SlideShape],
    charts_dir: Path,
    slide_height: Length,
    spacing: tuple[int, int] | None,
) -> None:
    """Add title, text box, and picture shapes in definition order."""
    body_placeholders = [
        cast(SlidePlaceholder, shape)
        for shape in slide.placeholders
        if shape.placeholder_format.type == PP_PLACEHOLDER.BODY
    ]
    for shape in shapes:
        if shape["type"] == "title":
            _add_title(slide, shape)
        elif shape["type"] == "text_box":
            _add_text_box(slide, shape, body_placeholders, spacing)
        else:
            _add_picture(slide, shape, charts_dir, slide_height)


def _add_title(slide: Slide, shape: TitleShape) -> None:
    """Set the layout title placeholder."""
    title = slide.shapes.title
    if title is None:
        raise ValueError("This layout has no title placeholder")
    title.text = shape["text"]
    if "size" in shape:
        for paragraph in title.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(shape["size"])
            _apply_paragraph_size(paragraph, shape["size"], bullet=False)
        title.text_frame.auto_size = MSO_AUTO_SIZE.NONE


def _add_text_box(
    slide: Slide,
    shape: TextBoxShape,
    body_placeholders: list[SlidePlaceholder],
    spacing: tuple[int, int] | None,
) -> None:
    """Fill the next body placeholder, or add a text box at an explicit position."""
    position = _inches_box(shape, "Text box")
    if position is None:
        if not body_placeholders:
            raise ValueError(
                "Text box has no position, and this layout has no unused body placeholder"
            )
        text_frame = body_placeholders.pop(0).text_frame
    else:
        left, top, width, height = position
        text_frame = slide.shapes.add_textbox(left, top, width, height).text_frame
        spacing = None
    box_size = shape["size"] if "size" in shape else None
    _write_paragraphs(text_frame, shape["paragraphs"], spacing, box_size)


def _template_body_spacing(slides: list[Slide]) -> tuple[int, int] | None:
    """Return the space before and after a template body paragraph, in EMU.

    The layout leaves that gap at zero. The template's body slides set it on
    each paragraph, and a newly added slide does not keep those paragraphs.
    """
    before: list[int] = []
    after: list[int] = []
    for slide in slides:
        for shape in slide.placeholders:
            if shape.placeholder_format.type != PP_PLACEHOLDER.BODY:
                continue
            body = cast(SlidePlaceholder, shape)
            for paragraph in body.text_frame.paragraphs:
                if not paragraph.text:
                    continue
                if paragraph.space_before is not None:
                    before.append(int(paragraph.space_before))
                if paragraph.space_after is not None:
                    after.append(int(paragraph.space_after))
    if not before and not after:
        return None
    return _most_common(before), _most_common(after)


def _most_common(values: list[int]) -> int:
    """Return the most frequent value, preferring the one that appears first."""
    if not values:
        return 0
    counts: dict[int, int] = {}
    for value in values:
        counts[value] = counts.get(value, 0) + 1
    return max(counts, key=lambda value: (counts[value], -values.index(value)))


def _write_paragraphs(
    text_frame: TextFrame,
    paragraphs: list[TextParagraph],
    spacing: tuple[int, int] | None = None,
    box_size: float | None = None,
) -> None:
    """Write paragraphs, applying bullets and run styling."""
    text_frame.clear()
    text_frame.word_wrap = True
    sized = False
    for index, spec in enumerate(paragraphs):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        level = spec["level"] if "level" in spec else 0
        bullet = bool(spec.get("bullet"))
        if bullet:
            _apply_bullet(paragraph, level)
        else:
            if level:
                paragraph.level = level
            _suppress_bullet(paragraph, level)
        if spacing is not None:
            paragraph.space_before = Emu(spacing[0])
            paragraph.space_after = Emu(spacing[1])
        paragraph_size = spec["size"] if "size" in spec else box_size
        applied_sizes: list[float] = []
        for run_spec in spec["runs"]:
            run = paragraph.add_run()
            run.text = run_spec["text"]
            size = run_spec["size"] if "size" in run_spec else paragraph_size
            if size is not None:
                run.font.size = Pt(size)
                sized = True
                applied_sizes.append(size)
            if run_spec.get("bold"):
                run.font.bold = True
            if run_spec.get("italic"):
                run.font.italic = True
            if "hyperlink" in run_spec:
                run.hyperlink.address = run_spec["hyperlink"]
        if paragraph_size is not None or applied_sizes:
            _apply_paragraph_size(
                paragraph,
                paragraph_size if paragraph_size is not None else applied_sizes[-1],
                bullet,
            )
    if sized:
        text_frame.auto_size = MSO_AUTO_SIZE.NONE


def _apply_paragraph_size(paragraph: _Paragraph, points: float, bullet: bool) -> None:
    """Record a point size where PowerPoint reads it for placeholder text.

    The layout autosizes body text and sizes bullets separately from the run.
    A run size on its own is scaled away unless autofit is off and the
    paragraph mark and bullet use the same size.
    """
    hundredths = int(round(points * 100))
    paragraph.font.size = Pt(points)
    paragraph._p.get_or_add_endParaRPr().sz = hundredths
    if not bullet:
        return
    properties = paragraph._p.get_or_add_pPr()
    for child in list(properties.findall(qn("a:buSzPts"))):
        properties.remove(child)
    bullet_size = OxmlElement("a:buSzPts")
    bullet_size.set("val", str(hundredths))
    bullet_font = properties.find(qn("a:buFont"))
    if bullet_font is None:
        _insert_before_default_run(properties, bullet_size)
    else:
        bullet_font.addprevious(bullet_size)


def _apply_bullet(paragraph: _Paragraph, level: int) -> None:
    """Show a bullet and indent the paragraph to match the template list style."""
    paragraph.level = level
    properties = paragraph._p.get_or_add_pPr()
    _remove_bullet_elements(properties)
    step = 457200
    properties.set("marL", str(457200 + level * step))
    properties.set("indent", str(-342900))
    bullet_font = OxmlElement("a:buFont")
    bullet_font.set("typeface", "Arial")
    bullet_character = OxmlElement("a:buChar")
    bullet_character.set("char", "●")
    _insert_before_default_run(properties, bullet_font)
    _insert_before_default_run(properties, bullet_character)


def _suppress_bullet(paragraph: _Paragraph, level: int) -> None:
    """Stop a body placeholder from inheriting its layout bullet and hanging indent.

    The layout indents wrapped lines further than the first line. A paragraph
    without a bullet keeps every line on the same left edge.
    """
    properties = paragraph._p.get_or_add_pPr()
    _remove_bullet_elements(properties)
    properties.set("marL", str(level * 457200))
    properties.set("indent", "0")
    _insert_before_default_run(properties, OxmlElement("a:buNone"))


def _remove_bullet_elements(properties: CT_TextParagraphProperties) -> None:
    """Remove an existing bullet choice so a new one can be written."""
    for tag in ("a:buFont", "a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip"):
        for child in list(properties.findall(qn(tag))):
            properties.remove(child)


def _insert_before_default_run(
    properties: CT_TextParagraphProperties,
    element: BaseOxmlElement,
) -> None:
    """Keep bullet elements ahead of the default run properties."""
    default_run = properties.find(qn("a:defRPr"))
    if default_run is None:
        properties.append(element)
    else:
        default_run.addprevious(element)


def _add_picture(
    slide: Slide,
    shape: PictureShape,
    charts_dir: Path,
    slide_height: Length,
) -> None:
    """Place a PNG, fitting it under the title when no position is given."""
    image_path = Path(shape["file"])
    if not image_path.is_absolute():
        image_path = charts_dir / image_path
    if not image_path.exists():
        raise FileNotFoundError(f"Chart not found: {image_path}")

    position = _inches_box(shape, "Picture")
    if position is None:
        _fit_picture_below_title(slide, image_path, slide_height)
        return

    left, top, width, height = position
    slide.shapes.add_picture(str(image_path), left, top, width, height)


def _inches_box(
    shape: TextBoxShape | PictureShape,
    kind: str,
) -> tuple[Length, Length, Length, Length] | None:
    """Return a positioned box, or None when the shape leaves placement to the layout."""
    keys = ("left", "top", "width", "height")
    present = [key for key in keys if key in shape]
    if not present:
        return None
    if len(present) != len(keys):
        missing = [key for key in keys if key not in shape]
        raise ValueError(f"{kind} position needs left, top, width, and height; missing {missing}")
    positioned = cast(_Positioned, shape)
    return (
        Inches(positioned["left"]),
        Inches(positioned["top"]),
        Inches(positioned["width"]),
        Inches(positioned["height"]),
    )


def _fit_picture_below_title(slide: Slide, image_path: Path, slide_height: Length) -> None:
    """Fit a chart in the area below the slide title, keeping its aspect ratio."""
    title = slide.shapes.title
    if title is None:
        raise ValueError("Cannot place a picture without a position on a slide that has no title")

    box_left = title.left
    box_top = title.top + title.height + Inches(0.05)
    box_width = title.width
    bottom_gap = Inches(PICTURE_BOTTOM_GAP_INCHES)
    box_height = slide_height - box_top - bottom_gap
    if box_height <= 0:
        raise ValueError(
            f"Picture bottom gap ({PICTURE_BOTTOM_GAP_INCHES} in) leaves no room under the title"
        )

    with Image.open(image_path) as image:
        image_width, image_height = image.size
    aspect = image_width / image_height

    if box_width / box_height > aspect:
        height = Emu(box_height)
        width = Emu(int(height * aspect))
    else:
        width = Emu(box_width)
        height = Emu(int(width / aspect))

    left = Emu(box_left + (box_width - width) // 2)
    top = Emu(box_top + (box_height - height) // 2)
    slide.shapes.add_picture(str(image_path), left, top, width, height)


def _reorder_slides(prs: PresentationDocument, order: list[Slide]) -> None:
    """Put the presentation slides into `order`. python-pptx has no public method for this."""
    slide_ids = prs.slides._sldIdLst
    entries = list(slide_ids.sldId_lst)
    entry_for_part = {
        id(prs.part.related_slide(entry.rId).part): entry
        for entry in entries
    }
    for entry in entries:
        slide_ids.remove(entry)
    for slide in order:
        slide_ids.append(entry_for_part[id(slide.part)])


def _delete_slide(prs: PresentationDocument, index: int) -> None:
    """Drop a slide. python-pptx has no public method for this."""
    slide_ids = prs.slides._sldIdLst
    slide_id = slide_ids.sldId_lst[index]
    prs.part.drop_rel(slide_id.rId)
    slide_ids.remove(slide_id)


def _parse_deck_entry(item: object, entry: int, path: Path) -> list[SlideDefinition]:
    """Return the slides for one deck entry, loading a shared file when it uses include."""
    label = f"{path.name} entry {entry}"
    fields = _object(item, label)
    if "include" in fields:
        return _parse_include(fields, label, path)
    return [_parse_slide(fields, entry, path)]


def _parse_include(fields: dict[str, object], label: str, deck_path: Path) -> list[SlideDefinition]:
    """Load one shared slide, or consecutive slides, and number them from this report."""
    _check_keys(fields, {"number", "include"}, set(), label)
    number = _slide_number(fields["number"], label)
    include_name = fields["include"]
    if not isinstance(include_name, str) or not include_name:
        raise ValueError(f"{label} include must be a file name")

    include_path = Path(include_name)
    if not include_path.is_absolute():
        include_path = deck_path.parent / include_path
    if not include_path.exists():
        raise FileNotFoundError(f"{label} include {include_name} does not exist.")

    with include_path.open() as include_file:
        raw = json.load(include_file)
    bodies = raw if isinstance(raw, list) else [raw]
    if not isinstance(raw, (dict, list)) or not bodies:
        raise ValueError(f"{include_path.name} must contain a slide or a list of slides.")

    slides: list[SlideDefinition] = []
    for offset, body in enumerate(bodies):
        body_fields = _object(body, f"{include_path.name} entry {offset + 1}")
        if "number" in body_fields or "include" in body_fields:
            raise ValueError(
                f"{include_path.name} must contain layout and shapes. "
                "Set the slide number on the include in the report."
            )
        body_fields["number"] = number + offset
        slides.append(_parse_slide(body_fields, offset + 1, include_path))
    return slides


def _parse_slide(fields: dict[str, object], entry: int, path: Path) -> SlideDefinition:
    """Return one validated slide, using the deck entry number in error messages."""
    label = f"{path.name} entry {entry}"
    _check_keys(fields, {"number", "layout", "shapes"}, set(), label)

    number = _slide_number(fields["number"], label)
    layout = fields["layout"]
    if not isinstance(layout, str) or not layout:
        raise ValueError(f"{label} layout must be a non-empty string")
    shapes = fields["shapes"]
    if not isinstance(shapes, list):
        raise ValueError(f"{label} shapes must be a list")

    return {
        "number": number,
        "layout": layout,
        "shapes": [
            _parse_shape(shape, entry, index, path)
            for index, shape in enumerate(shapes, start=1)
        ],
    }


def _slide_number(value: object, label: str) -> int:
    """Return a 1-based slide number."""
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise ValueError(f"{label} number must be an integer of 1 or greater")
    return value


def _parse_shape(item: object, entry: int, index: int, path: Path) -> SlideShape:
    """Return one validated title, text box, or picture."""
    label = f"{path.name} entry {entry} shape {index}"
    fields = _object(item, label)
    shape_type = fields.get("type")
    if shape_type == "title":
        _check_keys(fields, {"type", "text"}, {"size"}, label)
        title: TitleShape = {"type": "title", "text": _text(fields["text"], label)}
        if "size" in fields:
            title["size"] = _font_size(fields["size"], label)
        return title
    if shape_type == "text_box":
        _check_keys(fields, {"type"}, set(_POSITION_KEYS) | {"text", "paragraphs", "size"}, label)
        if ("text" in fields) == ("paragraphs" in fields):
            raise ValueError(f"{label} must have either text or paragraphs")
        paragraphs = (
            [_plain_paragraph(fields["text"], label)]
            if "text" in fields
            else _parse_paragraphs(fields["paragraphs"], label)
        )
        text_box: TextBoxShape = {"type": "text_box", "paragraphs": paragraphs}
        if "size" in fields:
            text_box["size"] = _font_size(fields["size"], label)
        _apply_position(text_box, fields, label)
        return text_box
    if shape_type == "picture":
        _check_keys(fields, {"type", "file"}, set(_POSITION_KEYS), label)
        file_name = fields["file"]
        if not isinstance(file_name, str) or not file_name:
            raise ValueError(f"{label} file must be a non-empty string")
        picture: PictureShape = {"type": "picture", "file": file_name}
        _apply_position(picture, fields, label)
        return picture
    raise ValueError(f"{label} type must be 'title', 'text_box', or 'picture'")


def _apply_position(shape: TextBoxShape | PictureShape, fields: dict[str, object], label: str) -> None:
    """Copy a complete box onto a shape. Leave the shape unchanged when it has no box."""
    present = [key for key in _POSITION_KEYS if key in fields]
    if not present:
        return
    if len(present) != len(_POSITION_KEYS):
        missing = [key for key in _POSITION_KEYS if key not in fields]
        raise ValueError(f"{label} needs left, top, width, and height; missing {missing}")
    for key in _POSITION_KEYS:
        value = fields[key]
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise ValueError(f"{label} {key} must be a number")
        shape[key] = float(value)


def _object(item: object, label: str) -> dict[str, object]:
    """Return a JSON object, rejecting arrays and other JSON values."""
    if not isinstance(item, dict):
        raise ValueError(f"{label} must be an object")
    return {str(key): value for key, value in item.items()}


def _font_size(value: object, label: str) -> float:
    """Return a font size in points. Deck files write the size as a string."""
    if isinstance(value, str):
        try:
            value = float(value)
        except ValueError:
            value = None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value <= 0:
        raise ValueError(f"{label} size must be a font size in points")
    return float(value)


def _text(value: object, label: str) -> str:
    """Return slide text, rejecting empty strings and non-strings."""
    if not isinstance(value, str) or not value:
        raise ValueError(f"{label} text must be a non-empty string")
    return value


def _plain_paragraph(value: object, label: str) -> TextParagraph:
    """Wrap a plain string as one paragraph with one run."""
    return {"runs": [{"text": _text(value, label)}]}


def _parse_paragraphs(value: object, label: str) -> list[TextParagraph]:
    """Return the paragraphs for a text box."""
    if not isinstance(value, list) or not value:
        raise ValueError(f"{label} paragraphs must be a non-empty list")
    return [
        _parse_paragraph(item, label, index)
        for index, item in enumerate(value, start=1)
    ]


def _parse_paragraph(item: object, parent: str, index: int) -> TextParagraph:
    """Return one paragraph, made from either text or styled runs."""
    label = f"{parent} paragraph {index}"
    fields = _object(item, label)
    _check_keys(fields, set(), {"text", "runs", "bullet", "level", "size"}, label)
    if ("text" in fields) == ("runs" in fields):
        raise ValueError(f"{label} must have either text or runs")
    paragraph: TextParagraph = {
        "runs": (
            [{"text": _text(fields["text"], label)}]
            if "text" in fields
            else _parse_runs(fields["runs"], label)
        )
    }
    if "bullet" in fields:
        if not isinstance(fields["bullet"], bool):
            raise ValueError(f"{label} bullet must be true or false")
        paragraph["bullet"] = fields["bullet"]
    if "level" in fields:
        level = fields["level"]
        if isinstance(level, bool) or not isinstance(level, int) or not 0 <= level <= 8:
            raise ValueError(f"{label} level must be an integer from 0 to 8")
        paragraph["level"] = level
    if "size" in fields:
        paragraph["size"] = _font_size(fields["size"], label)
    return paragraph


def _parse_runs(value: object, label: str) -> list[TextRun]:
    """Return the runs in one paragraph."""
    if not isinstance(value, list) or not value:
        raise ValueError(f"{label} runs must be a non-empty list")
    return [_parse_run(item, label, index) for index, item in enumerate(value, start=1)]


def _parse_run(item: object, parent: str, index: int) -> TextRun:
    """Return one run, including bold, italic, and hyperlink when present."""
    label = f"{parent} run {index}"
    fields = _object(item, label)
    _check_keys(fields, {"text"}, {"bold", "italic", "hyperlink", "size"}, label)
    run: TextRun = {"text": _text(fields["text"], label)}
    bold = fields.get("bold")
    if bold is not None:
        if not isinstance(bold, bool):
            raise ValueError(f"{label} bold must be true or false")
        run["bold"] = bold
    italic = fields.get("italic")
    if italic is not None:
        if not isinstance(italic, bool):
            raise ValueError(f"{label} italic must be true or false")
        run["italic"] = italic
    if "hyperlink" in fields:
        link = fields["hyperlink"]
        if not isinstance(link, str) or not link:
            raise ValueError(f"{label} hyperlink must be a non-empty string")
        run["hyperlink"] = link
    if "size" in fields:
        run["size"] = _font_size(fields["size"], label)
    return run


def _check_keys(
    fields: dict[str, object],
    required: set[str],
    optional: set[str],
    label: str,
) -> None:
    """Raise when required keys are absent or unknown keys are present."""
    keys = set(fields)
    missing = sorted(required - keys)
    unknown = sorted(keys - required - optional)
    if missing:
        raise ValueError(f"{label} is missing {missing}")
    if unknown:
        raise ValueError(f"{label} has unknown keys {unknown}")


def _resolve_deck(deck: str | Path | None, project_root: Path) -> Path:
    """Resolve a deck path from the working directory, then from the project root."""
    if deck is None:
        return project_root / DEFAULT_DECK
    path = Path(deck)
    if path.is_absolute() or path.exists():
        return path
    return project_root / path


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    """Parse the optional slide-deck path."""
    parser = argparse.ArgumentParser(description="Build a report deck from a slide JSON file.")
    parser.add_argument(
        "deck",
        nargs="?",
        help=f"Path to a slide deck JSON file. Defaults to {DEFAULT_DECK}.",
    )
    return parser.parse_args(argv)


def main(deck: str | Path | None = None) -> None:
    """Build the slides in a deck JSON file and save the presentation."""
    config = get_config()
    deck_path = _resolve_deck(deck, config.project_root)
    slides = load_slides(deck_path)
    output_path = config.report_dir / f"{deck_path.stem}.pptx"
    prs = Presentation(str(config.templates_dir / "report_slides_template.pptx"))

    logger.info(f"Loaded {len(slides)} slides from {deck_path}")
    insert_slides(prs, slides, config.charts_dir)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    logger.info(f"Saved report to {output_path}")


if __name__ == "__main__":
    main(_parse_args().deck)
